import psycopg2
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Conexión a PostgreSQL
conn = psycopg2.connect(
    host="localhost",
    port=5432,
    database="fashionstore",
    user="postgres",
    password="09091991"
)
conn.autocommit = True
cur = conn.cursor()

# Lista de tablas en orden de creación (respetando dependencias)
TABLAS = [
    "ciudades",
    "sucursales",
    "usuarios",
    "categorias",
    "colores",
    "tallas",
    "temporadas",
    "colecciones",
    "proveedores",
    "productos",
    "producto_variantes",
    "inventario",
    "movimientos_inventario",
    "reservas",
    "reserva_items",
    "pedidos",
    "pedido_items",
    "pagos",
]

# Estilos
HEADER_FILL = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
TITLE_FONT = Font(bold=True, size=14, color="1F4E79")
HEADER_BORDER = Border(
    bottom=Side(style="thin"),
    top=Side(style="thin"),
    left=Side(style="thin"),
    right=Side(style="thin"),
)
thin_border = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)
row_fill_alt = PatternFill(start_color="EAF1F8", end_color="EAF1F8", fill_type="solid")

wb = Workbook()

# ---------- HOJA ÍNDICE ----------
ws_idx = wb.active
ws_idx.title = "Índice"
ws_idx["A1"] = "FASHIONSTORE - BASE DE DATOS POBLADA"
ws_idx["A1"].font = Font(bold=True, size=16, color="1F4E79")
ws_idx["A2"] = "Plataforma Inteligente de Comercio Electrónico - Sistemas II"
ws_idx["A3"] = "Datos de Bolivia - Sucursal principal: Santa Cruz de la Sierra"
ws_idx["A4"] = ""

ws_idx["A6"] = "Contenido de hojas (tablas pobladas):"
ws_idx["A6"].font = Font(bold=True, size=12)

for i, t in enumerate(TABLAS, start=7):
    ws_idx.cell(row=i, column=1, value=f"- {t}")

ws_idx["A30"] = "Fecha de generación: 2026-09-02"
ws_idx["A31"] = "Base de datos: PostgreSQL 18"
ws_idx.column_dimensions['A'].width = 70

# ---------- HOJAS DE TABLAS ----------
for tabla in TABLAS:
    cur.execute(f"SELECT * FROM {tabla} LIMIT 0")
    colnames = [desc[0] for desc in cur.description]

    cur.execute(f"SELECT * FROM {tabla} ORDER BY 1")
    rows = cur.fetchall()

    ws = wb.create_sheet(title=tabla[:31])

    # Título
    ws["A1"] = f"TABLA: {tabla.upper()}  ({len(rows)} registros)"
    ws["A1"].font = TITLE_FONT

    # Encabezados en fila 3
    for c, col in enumerate(colnames, start=1):
        cell = ws.cell(row=3, column=c, value=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = HEADER_BORDER

    # Datos desde fila 4
    for r, row in enumerate(rows, start=4):
        for c, val in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=val)
            cell.border = thin_border
            if isinstance(val, str):
                cell.alignment = Alignment(vertical="center")
            else:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            if (r % 2) == 1:
                cell.fill = row_fill_alt

    # Autoancho
    for c, col in enumerate(colnames, start=1):
        max_len = len(col)
        for r in range(4, 4 + len(rows)):
            v = ws.cell(row=r, column=c).value
            if v is not None:
                max_len = max(max_len, len(str(v)))
        ws.column_dimensions[get_column_letter(c)].width = min(max_len + 3, 45)

cur.close()
conn.close()

out = r"D:\Primer parcial de Si2\E-commerce\database\fashionstore_poblacion.xlsx"
try:
    wb.save(out)
except PermissionError:
    out = r"D:\Primer parcial de Si2\E-commerce\database\fashionstore_poblacion_nuevo.xlsx"
    wb.save(out)
print(f"Archivo generado: {out}")
print(f"Archivo generado: {out}")
